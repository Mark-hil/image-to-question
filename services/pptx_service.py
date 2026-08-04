import os
from typing import Optional, Dict, Any, List
from pptx import Presentation
from services.pdf_service import parse_page_range

def extract_pptx_slides(file_path: str, page_range: Optional[str] = None) -> Dict[str, Any]:
    """
    Extract text content slide-by-slide from a PowerPoint (.pptx) file.
    Supports slide-range filtering (e.g., '1-5').
    """
    try:
        prs = Presentation(file_path)
        total_slides = len(prs.slides)
        
        if total_slides == 0:
            return {
                "text": "",
                "total_slides": 0,
                "processed_slides": 0
            }

        start_idx, end_idx = parse_page_range(page_range or "", total_slides)
        
        extracted_text_blocks: List[str] = []
        slide_count = 0

        for idx in range(start_idx, end_idx):
            if idx >= total_slides:
                break
            
            slide = prs.slides[idx]
            slide_number = idx + 1
            slide_lines: List[str] = []

            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_lines.append(shape.text.strip())

            if slide_lines:
                slide_content = "\n".join(slide_lines)
                extracted_text_blocks.append(f"--- [Slide {slide_number} of {total_slides}] ---\n{slide_content}")
                slide_count += 1

        full_text = "\n\n".join(extracted_text_blocks)
        return {
            "text": full_text,
            "total_slides": total_slides,
            "processed_slides": slide_count,
            "slide_range_used": f"{start_idx + 1}-{end_idx}"
        }
    except Exception as e:
        return {
            "text": "",
            "total_slides": 0,
            "processed_slides": 0,
            "error": str(e)
        }

def inspect_pptx_slides(file_path: str) -> Dict[str, Any]:
    """
    Inspect PowerPoint (.pptx) file and return total slide count and slide title pills.
    """
    try:
        prs = Presentation(file_path)
        total_slides = len(prs.slides)
        chapters = []

        for idx, slide in enumerate(prs.slides):
            slide_num = idx + 1
            slide_title = f"Slide {slide_num}"
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    first_line = shape.text.strip().splitlines()[0]
                    if len(first_line) > 30:
                        first_line = first_line[:30] + "..."
                    slide_title = f"Slide {slide_num}: {first_line}"
                    break
            
            chapters.append({
                "title": slide_title,
                "level": 1,
                "start_page": slide_num,
                "end_page": slide_num,
                "range": f"{slide_num}"
            })

        return {
            "total_pages": total_slides,
            "chapters": chapters
        }
    except Exception as e:
        return {
            "total_pages": 0,
            "chapters": [],
            "error": str(e)
        }
